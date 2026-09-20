#!/usr/bin/env python3
"""End-to-end self-test against a mock portal.

Proves the pipeline works before the real portal is calibrated:
  cycle 1 -> inserts
  cycle 2 -> same data, everything deduplicated
  cycle 3 -> one deadline moved, one new tender -> update + insert
Then exports to a sample Excel template and runs forecasting.

Run:  python selftest.py
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from pharmatender.db import Database
from pharmatender.excel_export import export, inspect_template
from pharmatender.forecasting import print_forecasts, run_forecasting
from pharmatender.pipeline import Screener
from pharmatender.sources.base import Attachment, BaseAdapter, ScrapedTender

WORK = Path("selftest_workspace")

CONFIG = {
    "database": str(WORK / "test.db"),
    "document_dir": str(WORK / "documents"),
    "download_documents": False,
    "delay_seconds": 0,
    "portfolio": ["trastuzumab", "pembrolizumab"],
    "sources": {"mock": {"departments": [{"label": "Medical Store"}]}},
}

# ------------------------------------------------------------------ mock site

CYCLE_A = [
    dict(tender_number="MOH/2026/0101", tender_title="Supply of Trastuzumab 440mg vials",
         publication_date="2026-01-12", submission_deadline="2026-02-10"),
    dict(tender_number="MOH/2026/0102", tender_title="Supply of Osimertinib 80mg tablets",
         publication_date="2026-01-14", submission_deadline="2026-02-12"),
    dict(tender_number="MOH/2026/0103", tender_title="Construction of staff car park",
         publication_date="2026-01-15", submission_deadline="2026-02-15"),
    dict(tender_number="MOH/2026/0104", tender_title="Supply of Meropenem 1g vials",
         publication_date="2026-01-16", submission_deadline="2026-02-18"),
    dict(tender_number="MOH/2026/0105", tender_title="توريد أدوية الأورام والعلاج الكيماوي",
         publication_date="2026-01-18", submission_deadline="2026-02-20"),
]

# same tenders, one deadline extended, plus a genuinely new one
CYCLE_C = [dict(t) for t in CYCLE_A]
CYCLE_C[0]["submission_deadline"] = "2026-02-24"      # extended
CYCLE_C.append(dict(tender_number="MOH/2026/0106",
                    tender_title="Supply of Pembrolizumab 100mg vial",
                    publication_date="2026-01-22",
                    submission_deadline="2026-02-28"))


class MockAdapter(BaseAdapter):
    id = "mock"
    name = "Mock portal"
    base_url = "https://example.invalid/tenders"

    def __init__(self, config, rows, logger=None):
        super().__init__(config, logger)
        self.rows = rows

    def list_tenders(self, run=None):
        for r in self.rows:
            yield ScrapedTender(
                issuing_authority="Ministry of Health - Kuwait",
                department="Medical Store",
                source_url=self.base_url,
                tender_page_url=f"{self.base_url}/{r['tender_number']}",
                currency="KWD",
                estimated_value=250000.0,
                raw_row=dict(r),
                **r,
            )

    def enrich(self, tender, run=None):
        return tender


# ------------------------------------------------------- sample Excel template


def make_template(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tender Tracker"
    ws["A1"] = "PHARMACEUTICAL TENDER TRACKER 2026"
    ws["A1"].font = Font(bold=True, size=14)
    headers = ["Tender Number", "Tender Title", "Issuing Authority",
               "Publication Date", "Submission Deadline", "Product Name",
               "Generic Name", "Dosage Form", "Strength", "Quantity",
               "Therapeutic Area", "Priority", "Oncology", "Estimated Value",
               "Currency", "Status", "Needs Manual Review", "Source URL", "Notes"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    ws.freeze_panes = "A4"
    dv = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("L4:L500")
    ws["D4"], ws["E4"] = "2026-01-01", "2026-02-01"
    ws["J4"] = 1000
    ws["N4"] = 250000
    for col, width in zip("ABCDEFGHIJKLMNOPQRS",
                          [16, 46, 26, 14, 16, 30, 22, 16, 14, 10, 20, 10, 10,
                           15, 9, 12, 18, 40, 24]):
        ws.column_dimensions[col].width = width
    wb.save(path)


# ------------------------------------------------------------------ run


def main() -> None:
    if WORK.exists():
        shutil.rmtree(WORK)
    WORK.mkdir(parents=True)

    template = WORK / "my_template.xlsx"
    make_template(template)
    print("### 1. INSPECT TEMPLATE\n")
    report = inspect_template(template, mapping_out=WORK / "mapping.yaml")
    sheet = report["sheets"][0]
    mapped = [c for c in sheet["columns"] if c["guessed_field"]]
    print(f"  header row detected: {sheet['header_row']}")
    print(f"  columns: {len(sheet['columns'])}, auto-mapped: {len(mapped)}, "
          f"unmapped: {len(sheet['columns']) - len(mapped)}")
    print(f"  dropdowns preserved: {len(sheet['data_validations'])}")
    for c in sheet["columns"][:6]:
        print(f"    {c['column']} {c['header']:<22} -> {c['guessed_field']}")

    db = Database(CONFIG["database"])

    print("\n### 2. SCREENING CYCLES\n")
    for label, rows in [("cycle 1 (first run)", CYCLE_A),
                        ("cycle 2 (nothing changed)", CYCLE_A),
                        ("cycle 3 (1 extended, 1 new)", CYCLE_C)]:
        stats = Screener(db, MockAdapter(CONFIG, rows), CONFIG).run()
        print(f"  {label:<28} seen={stats['listings_reviewed']} "
              f"new={stats['new_tenders']} updated={stats['updated_tenders']} "
              f"dupes={stats['duplicates_skipped']} onc={stats['oncology_tenders']} "
              f"review={stats['needs_review']}")

    print("\n### 3. STORED RESULTS\n")
    for r in db.query("SELECT tender_number, therapeutic_area, priority, "
                      "is_oncology, submission_deadline FROM tenders "
                      "ORDER BY tender_number"):
        print(f"  {r['tender_number']}  {r['therapeutic_area']:<22}"
              f"{r['priority']:<8}onc={r['is_oncology']}  "
              f"deadline={r['submission_deadline']}")

    ver = db.query("SELECT * FROM tender_versions")
    print(f"\n  change-history rows: {len(ver)}")
    for v in ver:
        print(f"    {v['field']}: {v['old_value']} -> {v['new_value']}")

    construction = db.query(
        "SELECT * FROM tenders WHERE tender_number = 'MOH/2026/0103'")
    print(f"  construction tender excluded from database: {not construction}")

    print("\n### 4. EXPORT TO TEMPLATE\n")
    mapping = yaml.safe_load((WORK / "mapping.yaml").read_text(encoding="utf-8"))
    out = export(db, mapping, WORK / "export.xlsx")
    wb = openpyxl.load_workbook(out)
    ws = wb["Tender Tracker"]
    print(f"  written: {out.name}")
    print(f"  header intact at row 3: {ws['A3'].value!r}")
    print(f"  freeze panes preserved: {ws.freeze_panes}")
    print(f"  dropdowns preserved: {len(ws.data_validations.dataValidation)}")
    for r in range(4, 4 + 6):
        if ws.cell(row=r, column=1).value:
            print(f"    row {r}: {ws.cell(row=r, column=1).value} | "
                  f"{str(ws.cell(row=r, column=2).value)[:34]:<34} | "
                  f"{ws.cell(row=r, column=11).value} | {ws.cell(row=r, column=12).value}")

    print("\n### 5. FORECASTING (deliberately thin history)\n")
    print_forecasts(run_forecasting(db, CONFIG), only_forecasts=False)


if __name__ == "__main__":
    main()
