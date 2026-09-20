"""The management report, as a formatted workbook.

The tracker export is a data drop: it fills the user's own template and is
meant to be worked in. This is the other half -- something that can be sent
to someone who will read it once, so it carries the branding, a summary page
and one sheet per question rather than a single wide grid.
"""
from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BRAND = "4D8B89"          # the Roots teal
BRAND_DARK = "2F5F5E"
INK = "1F2A30"
MUTED = "6B7C86"
BAND = "F2F7F7"           # alternating row tint
LINE = "D8E3E3"

TITLE_FONT = Font(name="Calibri", size=20, bold=True, color=BRAND_DARK)
SUB_FONT = Font(name="Calibri", size=10, color=MUTED)
H2_FONT = Font(name="Calibri", size=13, bold=True, color=BRAND_DARK)
HEAD_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Calibri", size=10, color=INK)
KPI_NUM = Font(name="Calibri", size=22, bold=True, color=BRAND_DARK)
KPI_LAB = Font(name="Calibri", size=9, bold=True, color=MUTED)

HEAD_FILL = PatternFill("solid", fgColor=BRAND)
KPI_FILL = PatternFill("solid", fgColor=BAND)
BAND_FILL = PatternFill("solid", fgColor=BAND)

THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet_title(ws, title: str, subtitle: str, width: int,
                 logo: Path | None = None) -> int:
    """Brand band across the top. Returns the first free row."""
    ws.row_dimensions[1].height = 46
    if logo and logo.is_file():
        try:
            img = XLImage(str(logo))
            img.height = 38
            img.width = int(38 * img.width / img.height) if img.height else 180
            ws.add_image(img, "A1")
        except Exception:                      # pillow missing or odd PNG
            ws["A1"] = "ROOTS PHARMACEUTICAL"
            ws["A1"].font = TITLE_FONT
    ws["A3"] = title
    ws["A3"].font = TITLE_FONT
    ws["A4"] = subtitle
    ws["A4"].font = SUB_FONT
    for col in range(1, width + 1):
        ws.cell(row=5, column=col).fill = PatternFill("solid", fgColor=BRAND)
    ws.row_dimensions[5].height = 3
    return 7


def _table(ws, row: int, headers: list[str], rows: list[list],
           widths: list[int] | None = None) -> int:
    """Write a banded table with a brand header. Returns the row after it."""
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BOX
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 26

    for i, data in enumerate(rows):
        r = row + 1 + i
        for j, v in enumerate(data, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.border = CELL_FONT, BOX
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if i % 2:
                c.fill = BAND_FILL

    for j, w in enumerate(widths or [18] * len(headers), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    if rows:
        ws.auto_filter.ref = (f"A{row}:{get_column_letter(len(headers))}"
                              f"{row + len(rows)}")
    return row + len(rows) + 2


def _kpis(ws, row: int, pairs: list[tuple[str, object]]) -> int:
    """A row of KPI tiles, two columns wide each."""
    col = 1
    for label, value in pairs:
        a = get_column_letter(col)
        b = get_column_letter(col + 1)
        ws.merge_cells(f"{a}{row}:{b}{row}")
        ws.merge_cells(f"{a}{row + 1}:{b}{row + 1}")
        v = ws[f"{a}{row}"]
        v.value = value
        v.font, v.fill = KPI_NUM, KPI_FILL
        v.alignment = Alignment(horizontal="center", vertical="center")
        l = ws[f"{a}{row + 1}"]
        l.value = label.upper()
        l.font, l.fill = KPI_LAB, KPI_FILL
        l.alignment = Alignment(horizontal="center", vertical="center")
        col += 3
    ws.row_dimensions[row].height = 34
    ws.row_dimensions[row + 1].height = 18
    return row + 3


def build_report(data: dict, logo: Path | None = None) -> bytes:
    """Render the report workbook. `data` comes from app.reports_data()."""
    wb = openpyxl.Workbook()
    today = data.get("as_of") or date.today().isoformat()
    scope = data.get("scope", "All tenders")

    # ---------------------------------------------------------- summary
    ws = wb.active
    ws.title = "Summary"
    r = _sheet_title(ws, "Tender Screening Report",
                     f"Kuwait MOH - Medical Store   |   as of {today}   |   {scope}",
                     10, logo)
    r = _kpis(ws, r, [("Tenders", data["tenders"]),
                      ("Line items", data["items"]),
                      ("Open tenders", data["open_tenders"]),
                      ("Roots SKU matches", data["roots_matches"])])
    r = _kpis(ws, r, [("Open + Roots match", data["open_roots"]),
                      ("MOH registered", data["moh_matches"]),
                      ("Needs review", data["needs_review"]),
                      ("Oncology", data["oncology"])])
    ws.cell(row=r, column=1, value="How to read this report").font = H2_FONT
    notes = [
        "Bidding opportunities lists open tenders whose screened molecule "
        "matches a Roots registered product.",
        "An 'exact' match means the molecule name agreed outright; 'verify' "
        "means it matched through transliteration or the brand bridge and "
        "should be checked against the tender document.",
        "Columns A-F of the tracker come from the MOH portal and the tender "
        "PDFs; G-H from the MOH price lists; I-K from the Roots "
        "registered-products sheet.",
    ]
    for i, n in enumerate(notes, start=1):
        c = ws.cell(row=r + i, column=1, value="• " + n)
        c.font = CELL_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r + i, start_column=1,
                       end_row=r + i, end_column=10)
        ws.row_dimensions[r + i].height = 30
    for j, w in enumerate([16] * 10, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # -------------------------------------------- bidding opportunities
    ws = wb.create_sheet("Bidding opportunities")
    r = _sheet_title(ws, "Bidding opportunities",
                     "Open tenders where a screened molecule matches a Roots SKU",
                     9, logo)
    _table(ws, r,
           ["Tender no.", "Closing date", "Item description", "Unit",
            "Tender qty", "Roots product", "Principal", "Registration",
            "Match"],
           data["opportunities"],
           [14, 14, 46, 9, 14, 30, 18, 15, 11])

    # ------------------------------------------------------ closing soon
    ws = wb.create_sheet("Closing soon")
    r = _sheet_title(ws, "Closing within 30 days",
                     "Every open tender in the window, matched or not", 5, logo)
    _table(ws, r,
           ["Tender no.", "Closing date", "Therapeutic area", "Title",
            "Priority"],
           data["closing_soon"], [14, 14, 22, 62, 12])

    # ------------------------------------------------------- the tracker
    ws = wb.create_sheet("Tracker")
    r = _sheet_title(ws, "Tracker", f"All line items in scope - {scope}", 11, logo)
    _table(ws, r,
           ["Tender no.", "Closing Date", "Item Description", "Theraputic Area",
            "Unit", "Tender QTY", "Registered products", "Company name",
            "Roots product name", "Principle name", "Registration status"],
           data["tracker"],
           [13, 13, 44, 20, 8, 13, 34, 26, 26, 18, 16])

    # --------------------------------------------------------- breakdowns
    ws = wb.create_sheet("Breakdowns")
    r = _sheet_title(ws, "Breakdowns", "Where the portfolio meets the pipeline",
                     4, logo)
    ws.cell(row=r, column=1, value="Roots SKU hits by principal").font = H2_FONT
    r = _table(ws, r + 1, ["Principal", "Line items"], data["by_principal"],
               [30, 14])
    ws.cell(row=r, column=1, value="Tenders by therapeutic area").font = H2_FONT
    _table(ws, r + 1, ["Therapeutic area", "Tenders"], data["by_area"],
           [30, 14])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
