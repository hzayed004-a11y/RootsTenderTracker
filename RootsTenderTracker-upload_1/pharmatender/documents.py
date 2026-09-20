"""Attachment handling: download, hash, parse, extract product line items.

Parsers degrade gracefully -- an unreadable attachment logs an error and the
cycle continues. Scanned PDFs with no text layer are flagged for OCR rather
than silently yielding nothing.
"""
from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from .normalize import (clean, extract_dosage_form, extract_pack_size,
                        extract_strength, fold, parse_amount, parse_quantity)

TEXT_EXT = {".txt", ".csv"}
PDF_EXT = {".pdf"}
WORD_EXT = {".docx", ".doc"}
EXCEL_EXT = {".xlsx", ".xlsm", ".xls"}
ZIP_EXT = {".zip", ".rar", ".7z"}
SUPPORTED = TEXT_EXT | PDF_EXT | WORD_EXT | EXCEL_EXT | ZIP_EXT


@dataclass
class ParsedDoc:
    filename: str
    text: str = ""
    tables: list[list[list[str]]] = None
    error: str = ""
    needs_ocr: bool = False

    def __post_init__(self):
        if self.tables is None:
            self.tables = []


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# ---------------------------------------------------------------- parsers


def _parse_pdf(data: bytes, name: str) -> ParsedDoc:
    doc = ParsedDoc(filename=name)
    try:
        import pdfplumber
    except ImportError:
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            doc.text = "\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception as exc:
            doc.error = f"pdf parse failed: {exc}"
            return doc
    else:
        try:
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                parts = []
                for page in pdf.pages:
                    parts.append(page.extract_text() or "")
                    for tbl in page.extract_tables() or []:
                        doc.tables.append(
                            [[clean(c) for c in row] for row in tbl if row]
                        )
                doc.text = "\n".join(parts)
        except Exception as exc:
            doc.error = f"pdf parse failed: {exc}"
            return doc

    if len(doc.text.strip()) < 40 and not doc.tables:
        doc.needs_ocr = True
        doc.error = "PDF has no extractable text layer (likely scanned) - OCR required"
    return doc


def _parse_docx(data: bytes, name: str) -> ParsedDoc:
    doc = ParsedDoc(filename=name)
    try:
        import docx
        d = docx.Document(io.BytesIO(data))
        doc.text = "\n".join(p.text for p in d.paragraphs)
        for tbl in d.tables:
            doc.tables.append([[clean(c.text) for c in row.cells] for row in tbl.rows])
    except Exception as exc:
        doc.error = f"docx parse failed: {exc}"
    return doc


def _parse_excel(data: bytes, name: str) -> ParsedDoc:
    doc = ParsedDoc(filename=name)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        parts = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [clean(c) for c in row]
                if any(cells):
                    rows.append(cells)
                    parts.append(" | ".join(cells))
            if rows:
                doc.tables.append(rows)
        doc.text = "\n".join(parts)
        wb.close()
    except Exception as exc:
        doc.error = f"excel parse failed: {exc}"
    return doc


def _parse_zip(data: bytes, name: str) -> list[ParsedDoc]:
    out: list[ParsedDoc] = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            for info in zf.infolist():
                if info.is_dir() or info.file_size > 60_000_000:
                    continue
                inner = f"{name}::{info.filename}"
                if Path(info.filename).suffix.lower() not in SUPPORTED:
                    out.append(ParsedDoc(filename=inner,
                                         error="unsupported file type inside archive"))
                    continue
                try:
                    out.extend(parse_bytes(zf.read(info), inner))
                except Exception as exc:
                    out.append(ParsedDoc(filename=inner, error=f"archive member: {exc}"))
    except zipfile.BadZipFile:
        out.append(ParsedDoc(filename=name,
                             error="not a readable ZIP (RAR/7z need external tooling)"))
    except Exception as exc:
        out.append(ParsedDoc(filename=name, error=f"zip parse failed: {exc}"))
    return out


def parse_bytes(data: bytes, filename: str) -> list[ParsedDoc]:
    ext = Path(filename.split("::")[-1]).suffix.lower()
    if ext in PDF_EXT:
        return [_parse_pdf(data, filename)]
    if ext in WORD_EXT:
        if ext == ".doc":
            return [ParsedDoc(filename=filename,
                              error="legacy .doc not supported - convert to .docx")]
        return [_parse_docx(data, filename)]
    if ext in EXCEL_EXT:
        if ext == ".xls":
            return [ParsedDoc(filename=filename,
                              error="legacy .xls not supported - convert to .xlsx")]
        return [_parse_excel(data, filename)]
    if ext in ZIP_EXT:
        return _parse_zip(data, filename)
    if ext in TEXT_EXT:
        return [ParsedDoc(filename=filename,
                          text=data.decode("utf-8", errors="replace"))]
    return [ParsedDoc(filename=filename, error=f"unsupported extension '{ext}'")]


# ------------------------------------------------- line-item extraction

ITEM_HEADER_HINTS = {
    "product": ["product", "item", "description", "drug", "medicine", "commodity",
                "صنف", "الصنف", "وصف", "دواء", "البيان"],
    "generic_name": ["generic", "inn", "scientific name", "active", "molecule",
                     "الاسم العلمي", "المادة الفعالة"],
    "brand_name": ["brand", "trade name", "الاسم التجاري"],
    "dosage_form": ["form", "dosage form", "presentation", "الشكل"],
    "strength": ["strength", "concentration", "التركيز"],
    "pack_size": ["pack", "packing", "pack size", "unit pack", "العبوة"],
    "quantity": ["quantity", "qty", "required qty", "annual quantity", "الكمية"],
    "unit": ["unit", "uom", "الوحدة"],
    "unit_price": ["unit price", "price", "rate", "السعر"],
    "item_no": ["item no", "s.n", "sr", "serial", "no.", "م", "رقم"],
}


# Longest hint first: "unit price" must win over "unit", otherwise a price
# column is misread as a unit-of-measure column and line values are lost.
_HINT_ORDER = sorted(
    ((fold(h), field) for field, hints in ITEM_HEADER_HINTS.items() for h in hints),
    key=lambda x: len(x[0]), reverse=True,
)


def _match_header(cell: str) -> str | None:
    f = fold(cell)
    if not f or len(f) > 60:
        return None
    for hint, field in _HINT_ORDER:
        if f == hint or hint in f:
            return field
    return None


def extract_items_from_tables(tables, source_name: str) -> list[dict]:
    """Find header rows in extracted tables and map columns to item fields."""
    items: list[dict] = []
    for table in tables:
        if len(table) < 2:
            continue
        header_idx, mapping = None, {}
        for i, row in enumerate(table[:8]):
            found = {}
            for j, cell in enumerate(row):
                f = _match_header(cell)
                if f and f not in found.values():
                    found[j] = f
            if len(found) >= 2 and "product" in found.values():
                header_idx, mapping = i, found
                break
        if header_idx is None:
            continue

        for row in table[header_idx + 1:]:
            rec: dict = {}
            for j, fieldname in mapping.items():
                if j < len(row):
                    val = clean(row[j])
                    if val:
                        rec[fieldname] = val
            name = rec.pop("product", None)
            if not name or len(name) < 3:
                continue
            item = {
                "product_name": name,
                "generic_name": rec.get("generic_name"),
                "brand_name": rec.get("brand_name"),
                "item_no": rec.get("item_no"),
                "dosage_form": rec.get("dosage_form") or extract_dosage_form(name),
                "strength": rec.get("strength") or extract_strength(name),
                "pack_size": rec.get("pack_size") or extract_pack_size(name),
                "extracted_from": source_name,
            }
            qty, unit = parse_quantity(rec.get("quantity"))
            item["quantity"] = qty
            item["unit"] = rec.get("unit") or unit
            price, _cur = parse_amount(rec.get("unit_price"))
            item["unit_price"] = price
            if qty and price:
                item["line_value"] = round(qty * price, 3)
            items.append(item)
    return items


_LINE_RX = re.compile(
    r"^\s*(?:(\d{1,3})[).\-\s]+)?([A-Za-z\u0600-\u06FF][\w\u0600-\u06FF\s\-/(),.]{4,80}?)"
    r"\s+(\d+(?:[.,]\d+)?\s*(?:mcg|mg|g|ml|iu|%)\b[^\n]{0,30})",
    re.MULTILINE,
)


def extract_items_from_text(text: str, source_name: str, limit: int = 300) -> list[dict]:
    """Fallback for documents whose product lists are prose, not tables."""
    items = []
    for m in _LINE_RX.finditer(text or ""):
        name = clean(f"{m.group(2)} {m.group(3)}")
        if len(name) < 6:
            continue
        items.append({
            "item_no": m.group(1),
            "product_name": name,
            "strength": extract_strength(name),
            "dosage_form": extract_dosage_form(name),
            "pack_size": extract_pack_size(name),
            "extracted_from": source_name,
            "needs_review": 1,
            "review_reason": "Parsed from unstructured text - verify against source",
        })
        if len(items) >= limit:
            break
    return items


def extract_items(parsed: ParsedDoc) -> list[dict]:
    items = extract_items_from_tables(parsed.tables, parsed.filename)
    if not items:
        items = extract_items_from_text(parsed.text, parsed.filename)
    return items
